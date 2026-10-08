"""Audit our city data against a licensed benchmark workbook (Student Source layout).

The workbook is QA-only input: nothing it holds may be published as our figure.

Usage: .venv/bin/python -I audit_vs_benchmark_pack.py <pack.xlsx> <out.json>
Reads the workbook (untrusted data) with openpyxl only; queries our DB read-only.
"""
import sys, re, json, statistics, datetime, collections

import os
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import openpyxl  # noqa: E402
from sqlalchemy import text  # noqa: E402
from app.database import SessionLocal  # noqa: E402
from app.scrapers.scheme_matching import norm_name, norm_pc, name_score, build_index, best_match  # noqa: E402

PACK, OUT = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(PACK, data_only=True, read_only=True)
db = SessionLocal()
out = {}


def s(v):
    return "" if v is None else str(v).strip()


def num(v):
    """int from 236, '236', '[209]' -> (236, bracketed)"""
    if v is None:
        return None, False
    if isinstance(v, (int, float)):
        return int(v), False
    t = str(v).strip()
    br = t.startswith("[")
    m = re.search(r"\d[\d,]*", t)
    return (int(m.group(0).replace(",", "")), br) if m else (None, br)


def year(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return int(v)
    m = re.search(r"(19|20)\d{2}", str(v))
    return int(m.group(0)) if m else None


def opnorm(n):
    t = re.sub(r"[^a-z0-9 ]", " ", (n or "").lower())
    drop = {"ltd", "limited", "plc", "the", "student", "students", "living", "accommodation", "group", "halls", "homes", "home", "property"}
    return {w for w in t.split() if w not in drop}


def op_same(a, b):
    A, B = opnorm(a), opnorm(b)
    if not A or not B:
        return False
    return len(A & B) / len(A | B) >= 0.5 or A <= B or B <= A


# ---------------- pack: private stock ----------------
ws = wb["Private Sector Stock"]
stock = []
for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
    if i == 1 or not s(row[2]):
        continue
    if s(row[1]) == "" and s(row[2]) == "":
        continue
    beds, br = num(row[4])
    if beds is None:
        continue
    stock.append({"row": i, "operator": s(row[1]), "name": s(row[2]), "postcode": s(row[3]),
                  "beds": beds, "bracketed": br, "build_raw": s(row[5]), "build": year(row[5]),
                  "capture": s(row[6]), "note": s(row[7])})
out["pack_stock_rows"] = len(stock)
out["pack_stock_beds_counted"] = sum(r["beds"] for r in stock if not r["bracketed"])
out["pack_stock_bracketed"] = [{"name": r["name"], "beds": r["beds"], "note": r["note"], "build": r["build"]} for r in stock if r["bracketed"]]

# ---------------- pack: university ----------------
ws = wb["University Stock"]
uni = []
for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
    if i == 1 or not s(row[2]) or s(row[2]).lower() == "total":
        continue
    beds, br = num(row[3])
    if beds is None:
        continue
    uni.append({"university": s(row[1]), "name": s(row[2]), "beds": beds, "bracketed": br, "note": s(row[4]) if len(row) > 4 else ""})
out["pack_uni_rows"] = len(uni)
out["pack_uni_beds_counted"] = sum(r["beds"] for r in uni if not r["bracketed"])

# ---------------- pack: pipeline ----------------
ws = wb["PBSA Pipeline"]
pipe = []
for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
    if i == 1 or s(row[0]).lower() != "birmingham":
        continue
    beds, _ = num(row[14])
    approved, _ = num(row[15])
    ref = s(row[4]).upper()
    m = re.search(r"(\d{4})/(\d{5})", ref)
    key = f"{m.group(1)}/{m.group(2)}" if m else None
    pipe.append({"row": i, "site": s(row[1]), "address": s(row[2]), "postcode": s(row[3]), "ref": s(row[4]), "key": key,
                 "status": s(row[7]), "stage": s(row[8]), "operator": s(row[9]), "developer": s(row[10]),
                 "sector": s(row[12]), "beds": beds, "approved": approved, "delivery": s(row[17]),
                 "opening_2026": num(row[18])[0]})
out["pack_pipeline_rows"] = len(pipe)
out["pack_pipeline_approved_beds"] = sum(r["approved"] or 0 for r in pipe)
out["pack_pipeline_2026_openings"] = sum(r["opening_2026"] or 0 for r in pipe)
out["pack_pipeline_application_beds"] = sum((r["beds"] or 0) for r in pipe if not r["status"].startswith("Planning Approved"))

# ---------------- pack: prices 26-27 ----------------
ws = wb["PBSA Prices 22-23 - 26-27"]
prices = []
for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
    if i == 1:
        continue
    p2627 = row[40]
    if not isinstance(p2627, (int, float)):
        continue
    prices.append({"name": s(row[46]) or s(row[0]), "operator": s(row[47]) or s(row[1]), "postcode": s(row[5]),
                   "beds": num(row[6])[0], "room_type": s(row[44]) or s(row[12]), "sub": s(row[45]) or s(row[13]),
                   "p2627": float(p2627), "weeks2627": num(row[41])[0],
                   "p2526": float(row[31]) if isinstance(row[31], (int, float)) else None,
                   "p2425": float(row[22]) if isinstance(row[22], (int, float)) else None,
                   "capture": s(row[48])})
out["pack_price_rows_2627"] = len(prices)
out["pack_price_room_types"] = collections.Counter(p["room_type"] for p in prices).most_common()

# ---------------- pack: HESA ----------------
ws = wb["HESA Numbers"]
hesa_pack = {}
years = None
for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
    if i == 1:
        years = [s(v) for v in row[2:10]]
        continue
    if i > 6:
        break
    name = s(row[1])
    hesa_pack[name] = {years[j]: int(row[2 + j]) for j in range(len(years)) if isinstance(row[2 + j], (int, float))}

# ---------------- ours ----------------
schemes = db.execute(text("""
 select s.id, s.name, s.postcode, oc.name as operator, s.beds_total, s.total_units, s.build_year, s.nominations,
        s.operating_status, s.source, s.field_provenance, s.scheme_type
 from existing_schemes s join councils c on c.id=s.council_id left join companies oc on oc.id=s.operator_company_id
 where c.name='Birmingham'""")).mappings().all()
priv = [r for r in schemes if r["scheme_type"] == "PBSA" and not (r["operator"] or "").lower().find("university") >= 0]
univ = [r for r in schemes if r["scheme_type"] == "PBSA" and (r["operator"] or "").lower().find("university") >= 0]
btr = [r for r in schemes if r["scheme_type"] != "PBSA"]
out["ours_private"] = {"schemes": len(priv), "beds": sum((r["beds_total"] or r["total_units"] or 0) for r in priv)}
out["ours_university"] = {"schemes": len(univ), "beds": sum((r["beds_total"] or 0) for r in univ)}
out["ours_btr_excluded"] = [r["name"] for r in btr]


class Obj:  # for build_index
    def __init__(self, r):
        self.id, self.name, self.postcode, self.r = r["id"], r["name"], r["postcode"], r


idx_priv = build_index([Obj(r) for r in priv])
idx_univ = build_index([Obj(r) for r in univ])

obs = db.execute(text("""
 select o.scheme_id, o.field, o.value_num, o.value_text, o.source, o.basis, o.source_reference
 from scheme_observations o join existing_schemes s on s.id=o.scheme_id join councils c on c.id=s.council_id
 where c.name='Birmingham'""")).mappings().all()
obs_by = collections.defaultdict(list)
for o in obs:
    obs_by[o["scheme_id"]].append(o)

# ---- census comparison ----
census_rows, used = [], set()
own_beds_checks, op_checks, build_checks = [], [], []
for r in stock:
    m, score = best_match(idx_priv, r["name"], r["postcode"], used=used)
    if m is None:
        census_rows.append({**r, "match": None})
        continue
    used.add(m.id)
    o = m.r
    prov = o["field_provenance"] or {}
    our_beds = o["beds_total"] or o["total_units"]
    beds_src = (prov.get("beds_total") or {}).get("source")
    op_src = (prov.get("operator") or {}).get("source")
    by_src = (prov.get("build_year") or {}).get("source")
    observed = (prov.get("observed") or {}).get("value") or []
    # own-source bed observations
    own_beds = [(x["source"], int(x["value_num"])) for x in obs_by[m.id] if x["field"] == "beds_total" and x["value_num"] is not None]
    for src, val in own_beds:
        own_beds_checks.append({"scheme": o["name"], "source": src, "ours": val, "pack": r["beds"],
                                "diff_pct": round(100 * (val - r["beds"]) / r["beds"], 1)})
    own_ops = {(x["source"], x["value_text"]) for x in obs_by[m.id] if x["field"] == "operator" and x["value_text"]}
    for src, val in own_ops:
        op_checks.append({"scheme": o["name"], "source": src, "ours": val, "pack": r["operator"], "same": op_same(val, r["operator"])})
    own_by = [(x["source"], int(x["value_num"]), x["basis"]) for x in obs_by[m.id] if x["field"] == "build_year" and x["value_num"] is not None]
    for src, val, basis in own_by:
        build_checks.append({"scheme": o["name"], "source": src, "basis": basis, "ours": val, "pack": r["build"], "pack_raw": r["build_raw"],
                             "same": (r["build"] is not None and val == r["build"]), "within1": (r["build"] is not None and abs(val - r["build"]) <= 1)})
    letting_obs = [x for x in obs_by[m.id] if x["field"] == "booking_status"]
    census_rows.append({**r, "match": o["name"], "score": score, "our_beds": our_beds, "beds_source": beds_src,
                        "our_operator": o["operator"], "operator_source": op_src, "op_same": op_same(o["operator"], r["operator"]),
                        "our_build": o["build_year"], "build_source": by_src,
                        "our_status": o["operating_status"], "nominations": o["nominations"],
                        "observed_sources": observed, "letting_checked": bool(letting_obs),
                        "letting_state": (letting_obs[0]["value_text"] if letting_obs else None)})
unmatched_pack = [r for r in census_rows if r["match"] is None]
matched = [r for r in census_rows if r["match"]]
out["census"] = {
    "pack_rows": len(stock), "matched": len(matched), "unmatched_pack_rows": [{"name": r["name"], "beds": r["beds"], "bracketed": r["bracketed"], "note": r["note"]} for r in unmatched_pack],
    "ours_not_in_pack": [o["name"] for o in priv if o["id"] not in used],
    "beds_equal": sum(1 for r in matched if r["our_beds"] == r["beds"]),
    "beds_source_counts": collections.Counter(r["beds_source"] for r in matched).most_common(),
    "operator_source_counts": collections.Counter(r["operator_source"] for r in matched).most_common(),
    "build_source_counts": collections.Counter(r["build_source"] for r in matched).most_common(),
    "operator_same": sum(1 for r in matched if r["op_same"]),
    "operator_diff": [{"scheme": r["name"], "ours": r["our_operator"], "pack": r["operator"]} for r in matched if not r["op_same"]],
    "build_equal": sum(1 for r in matched if r["build"] is not None and r["our_build"] == r["build"]),
    "build_pack_na": sum(1 for r in matched if r["build"] is None),
    "build_diff": [{"scheme": r["name"], "ours": r["our_build"], "pack": r["build_raw"]} for r in matched if r["build"] is not None and r["our_build"] != r["build"]],
    "observed_by_own_source": sum(1 for r in matched if r["observed_sources"]),
    "letting_checked": sum(1 for r in matched if r["letting_checked"]),
    "no_letting_pack": [r["name"] for r in matched if "no 26-27 letting" in r["note"].lower()],
    "no_letting_ours": [r["name"] for r in matched if r["our_status"] == "no_letting_presence"],
    "nominations_pack": [r["name"] for r in matched if "nomination" in r["note"].lower()],
    "nominations_ours": [r["name"] for r in matched if r["nominations"]],
    "own_beds_checks": own_beds_checks,
    "own_operator_checks": op_checks,
    "own_build_checks": build_checks,
    "sold_out_by_pack_beds": [{"scheme": r["name"], "beds": r["beds"], "state": r["letting_state"]} for r in matched if r["letting_state"]],
}

# ---- university ----
uni_rows, used_u = [], set()
for r in uni:
    m, score = best_match(idx_univ, r["name"], None, threshold=0.34, used=used_u)
    if m:
        used_u.add(m.id)
    uni_rows.append({**r, "match": m.r["name"] if m else None, "our_beds": m.r["beds_total"] if m else None, "our_source": m.r["source"] if m else None})
out["university"] = {"pack_rows": len(uni), "matched": sum(1 for r in uni_rows if r["match"]),
                     "beds_equal": sum(1 for r in uni_rows if r["match"] and r["our_beds"] == r["beds"]),
                     "unmatched": [{"name": r["name"], "beds": r["beds"], "bracketed": r["bracketed"], "note": r["note"]} for r in uni_rows if not r["match"]],
                     "ours_sources": collections.Counter(r["source"] for r in univ).most_common(),
                     "ours_not_in_pack": [o["name"] for o in univ if o["id"] not in used_u]}

# ---- pipeline ----
apps = db.execute(text("""
 select p.id, p.reference, p.status, p.decision_date, p.pbsa_beds, p.expected_delivery_year, p.construction_status,
        p.address, p.postcode, p.is_pbsa, left(p.description, 160) as description, p.submitted_date
 from planning_applications p join councils c on c.id=p.council_id
 where c.name='Birmingham' and (p.is_pbsa = true or p.description ilike '%student%')""")).mappings().all()
this_year = datetime.date.today().year
by_key = {}
for a in apps:
    m = re.search(r"(\d{4})/(\d{5})", a["reference"] or "")
    if m:
        by_key.setdefault(f"{m.group(1)}/{m.group(2)}", []).append(a)


def our_status(a):
    st = (a["status"] or "").lower()
    if st.startswith("approv"):
        if a["expected_delivery_year"] and a["expected_delivery_year"] < this_year and not a["construction_status"]:
            return "approved, overdue — no construction observed"
        return "approved"
    return st or "unknown"


pipe_rows = []
for r in pipe:
    ours = by_key.get(r["key"], []) if r["key"] else []
    alt = []
    if not ours and r["postcode"]:
        pc = norm_pc(r["postcode"])
        alt = [a for a in apps if norm_pc(a["postcode"]) == pc or (a["address"] and pc[:4] in norm_pc(a["address"]) and a["pbsa_beds"])]
    a = (ours or alt or [None])[0]
    pipe_rows.append({"site": r["site"], "ref": r["ref"], "pack_status": r["status"], "pack_beds": r["beds"], "pack_approved": r["approved"],
                      "pack_delivery": r["delivery"], "pack_opening_2026": r["opening_2026"], "sector": r["sector"],
                      "found_by": "ref" if ours else ("postcode" if alt else None),
                      "our_ref": a["reference"] if a else None, "our_status": our_status(a) if a else None,
                      "our_beds": a["pbsa_beds"] if a else None, "our_delivery": a["expected_delivery_year"] if a else None,
                      "our_decision": str(a["decision_date"]) if a and a["decision_date"] else None,
                      "beds_equal": (a is not None and a["pbsa_beds"] == (r["approved"] or r["beds"]))})
pack_keys = {r["key"] for r in pipe if r["key"]}
ours_big_not_in_pack = [{"ref": a["reference"], "status": our_status(a), "beds": a["pbsa_beds"], "delivery": a["expected_delivery_year"],
                         "address": (a["address"] or "")[:70], "decision": str(a["decision_date"])}
                        for a in apps if a["pbsa_beds"] and a["pbsa_beds"] >= 100 and (a["status"] or "").lower().startswith("approv")
                        and not any(k in (a["reference"] or "") for k in pack_keys)]
ours_big_not_in_pack.sort(key=lambda x: -x["beds"])
approved_apps = [a for a in apps if (a["status"] or "").lower().startswith("approv")]
out["pipeline"] = {
    "pack_rows": len(pipe), "pack_approved_rows": sum(1 for r in pipe if r["status"].startswith("Planning Approved")),
    "pack_approved_beds": out["pack_pipeline_approved_beds"], "pack_2026_openings": out["pack_pipeline_2026_openings"],
    "pack_application_beds": out["pack_pipeline_application_beds"],
    "ours_apps": len(apps), "ours_with_beds": sum(1 for a in apps if a["pbsa_beds"]),
    "ours_approved_beds": sum(a["pbsa_beds"] or 0 for a in approved_apps),
    "ours_overdue_beds": sum(a["pbsa_beds"] or 0 for a in approved_apps if our_status(a).startswith("approved, overdue")),
    "ours_dated_beds": {y: sum(a["pbsa_beds"] or 0 for a in approved_apps if a["expected_delivery_year"] == y) for y in (2026, 2027, 2028, 2029)},
    "found_by_ref": sum(1 for r in pipe_rows if r["found_by"] == "ref"), "found_by_postcode": sum(1 for r in pipe_rows if r["found_by"] == "postcode"),
    "not_found": [r["site"] + " " + r["ref"] for r in pipe_rows if not r["found_by"]],
    "beds_equal": sum(1 for r in pipe_rows if r["beds_equal"]),
    "beds_missing_ours": [{"site": r["site"], "ref": r["ref"], "pack_beds": r["pack_beds"], "our_ref": r["our_ref"], "our_status": r["our_status"]} for r in pipe_rows if r["found_by"] and r["our_beds"] is None],
    "beds_differ": [{"site": r["site"], "ref": r["ref"], "pack_beds": r["pack_beds"], "our_beds": r["our_beds"], "our_ref": r["our_ref"]} for r in pipe_rows if r["found_by"] and r["our_beds"] is not None and not r["beds_equal"]],
    "rows": pipe_rows,
    "ours_approved_100plus_not_in_pack": ours_big_not_in_pack,
}

# ---- rents ----
rents = db.execute(text("""
 select r.scheme_id, r.room_type, r.sub_classification, r.rent_per_week, r.academic_year, r.source, r.contract_length_weeks
 from scheme_rents r join existing_schemes s on s.id=r.scheme_id join councils c on c.id=s.council_id
 where c.name='Birmingham' and r.is_current and r.rent_per_week is not null""")).mappings().all()
name_by_id = {r["id"]: r["name"] for r in priv}
TYPE_MAP = {"studio": "Studio", "en-suite": "En-suite", "ensuite": "En-suite", "en suite": "En-suite",
            "non en-suite": "Non en-suite", "non-ensuite": "Non en-suite", "standard": "Non en-suite",
            "one bed apartment": "1-bed apartment", "1 bed apartment": "1-bed apartment", "one bed flat": "1-bed apartment",
            "two bed apartment": "2-bed apartment", "2 bed apartment": "2-bed apartment", "twodio": "Twodio"}


def tnorm(t):
    k = (t or "").strip().lower()
    return TYPE_MAP.get(k, (t or "").strip())


# pack per scheme/type
pack_pairs = collections.defaultdict(list)
pack_min = collections.defaultdict(list)
pack_growth = []
unmatched_price_names = collections.Counter()
used_p = {}
for p in prices:
    key = (p["name"], norm_pc(p["postcode"]))
    if key not in used_p:
        m, _ = best_match(idx_priv, p["name"], p["postcode"])
        used_p[key] = m.id if m else None
    sid = used_p[key]
    if sid is None:
        unmatched_price_names[p["name"]] += 1
        continue
    pack_pairs[(sid, tnorm(p["room_type"]))].append(p["p2627"])
    pack_min[sid].append(p["p2627"])
    if p["p2526"]:
        pack_growth.append(100 * (p["p2627"] - p["p2526"]) / p["p2526"])
our_pairs = collections.defaultdict(list)
our_from = collections.defaultdict(list)
for r in rents:
    t = tnorm(r["room_type"])
    if t in ("From (advertised)",):
        our_from[r["scheme_id"]].append((float(r["rent_per_week"]), r["source"], r["academic_year"]))
    elif t in ("Studio", "En-suite", "Non en-suite", "1-bed apartment", "2-bed apartment") and r["academic_year"] == "2026-27":
        our_pairs[(r["scheme_id"], t)].append(float(r["rent_per_week"]))
pair_rows = []
for k, vals in our_pairs.items():
    if k in pack_pairs:
        om, pm = statistics.median(vals), statistics.median(pack_pairs[k])
        pair_rows.append({"scheme": name_by_id.get(k[0]), "type": k[1], "ours_median": om, "ours_n": len(vals), "ours_min": min(vals), "ours_max": max(vals),
                          "pack_median": pm, "pack_n": len(pack_pairs[k]), "pack_min": min(pack_pairs[k]), "pack_max": max(pack_pairs[k]),
                          "diff_pct": round(100 * (om - pm) / pm, 1),
                          "range_overlap": not (max(vals) < min(pack_pairs[k]) or min(vals) > max(pack_pairs[k]))})
pair_rows.sort(key=lambda x: abs(x["diff_pct"]), reverse=True)
from_rows = []
for sid, vals in our_from.items():
    if sid in pack_min:
        ours = min(v for v, _, _ in vals)
        from_rows.append({"scheme": name_by_id.get(sid), "ours_from": ours, "sources": sorted({s_ for _, s_, _ in vals}),
                          "pack_min_2627": min(pack_min[sid]), "diff_pct": round(100 * (ours - min(pack_min[sid])) / min(pack_min[sid]), 1)})
from_rows.sort(key=lambda x: abs(x["diff_pct"]), reverse=True)
diffs = [r["diff_pct"] for r in pair_rows]
fdiffs = [r["diff_pct"] for r in from_rows]
out["rents"] = {
    "pack_rows_2627": len(prices), "pack_schemes_priced": len(pack_min), "pack_unmatched_names": unmatched_price_names.most_common(),
    "pack_median_2627": statistics.median([p["p2627"] for p in prices]),
    "pack_yoy_2526_to_2627_median_pct": round(statistics.median(pack_growth), 1) if pack_growth else None,
    "pack_yoy_n": len(pack_growth),
    "ours_rows": len(rents), "ours_schemes": len({r["scheme_id"] for r in rents}),
    "ours_2627_typed_pairs": len(our_pairs), "pairs_compared": len(pair_rows),
    "pairs_within_5": sum(1 for d in diffs if abs(d) <= 5), "pairs_within_10": sum(1 for d in diffs if abs(d) <= 10),
    "pairs_range_overlap": sum(1 for r in pair_rows if r["range_overlap"]),
    "pairs_median_diff_pct": round(statistics.median(diffs), 1) if diffs else None,
    "pair_rows": pair_rows,
    "from_compared": len(from_rows), "from_within_5": sum(1 for d in fdiffs if abs(d) <= 5), "from_within_10": sum(1 for d in fdiffs if abs(d) <= 10),
    "from_median_diff_pct": round(statistics.median(fdiffs), 1) if fdiffs else None, "from_rows": from_rows,
    "ours_type_medians": {t: statistics.median([float(r["rent_per_week"]) for r in rents if tnorm(r["room_type"]) == t]) for t in ("Studio", "En-suite", "Non en-suite")},
    "pack_type_medians": {t: statistics.median([p["p2627"] for p in prices if tnorm(p["room_type"]) == t]) for t in ("Studio", "En-suite", "Non en-suite") if any(tnorm(p["room_type"]) == t for p in prices)},
}

# ---- HESA ----
ours_hesa = db.execute(text("""
 select i.name, e.academic_year, e.full_time_students from hesa_enrolments e join institutions i on i.id=e.institution_id
 where i.city='Birmingham' and e.level='all' and e.source like 'hesa_table%'""")).all()
oh = collections.defaultdict(dict)
for n, y, v in ours_hesa:
    oh[n][y] = v
alias = {"Birmingham Newman University": "Newman University"}
hesa_rows, hesa_pack_total, hesa_our_total = [], collections.Counter(), collections.Counter()
for n, yrs in hesa_pack.items():
    on = alias.get(n, n)
    for y, v in yrs.items():
        ov = oh.get(on, {}).get(y)
        hesa_rows.append({"institution": n, "year": y, "pack": v, "ours": ov, "diff": (ov - v) if ov is not None else None})
        hesa_pack_total[y] += v
        if ov is not None:
            hesa_our_total[y] += ov
out["hesa"] = {"cells": len(hesa_rows), "equal": sum(1 for r in hesa_rows if r["diff"] == 0),
               "diff_rows": [r for r in hesa_rows if r["diff"] not in (0, None)],
               "totals": {y: {"pack": hesa_pack_total[y], "ours": hesa_our_total[y]} for y in sorted(hesa_pack_total)},
               "ours_2024_25": {n: oh[n].get("2024-25") for n in oh}}

# ---- summary ratios ----
out["summary"] = {"pack_private": out["pack_stock_beds_counted"], "pack_uni": out["pack_uni_beds_counted"],
                  "pack_ratio_current_inc_2026": round((out["pack_stock_beds_counted"] + out["pack_uni_beds_counted"] + out["pack_pipeline_2026_openings"]) / hesa_pack_total["2023-24"], 4)}
json.dump(out, open(OUT, "w"), indent=1, default=str)
print(json.dumps({k: v for k, v in out.items() if k not in ("census", "pipeline", "rents", "hesa", "university")}, indent=1, default=str))
c = out["census"]
print("CENSUS:", {k: v for k, v in c.items() if not isinstance(v, list)})
print(" unmatched pack:", c["unmatched_pack_rows"]); print(" ours not in pack:", c["ours_not_in_pack"])
print(" operator diffs:", c["operator_diff"]); print(" build diffs:", c["build_diff"])
print(" own beds checks:", c["own_beds_checks"]); print(" own operator checks (not same):", [x for x in c["own_operator_checks"] if not x["same"]], "of", len(c["own_operator_checks"]))
print(" own build checks:", c["own_build_checks"])
print(" no-letting pack/ours:", c["no_letting_pack"], c["no_letting_ours"]); print(" nominations:", c["nominations_pack"], c["nominations_ours"])
u = out["university"]; print("UNIVERSITY:", {k: v for k, v in u.items()})
p = out["pipeline"]; print("PIPELINE:", {k: v for k, v in p.items() if k not in ("rows",)})
for r in p["rows"]:
    print("  ", r["site"][:34].ljust(34), (r["ref"] or "")[:14].ljust(14), r["pack_status"][:22].ljust(22), "pack", r["pack_approved"] or r["pack_beds"], r["pack_delivery"][:16].ljust(16), "| ours", r["our_ref"], r["our_status"], r["our_beds"], r["our_delivery"], r["found_by"])
rr = out["rents"]; print("RENTS:", {k: v for k, v in rr.items() if k not in ("pair_rows", "from_rows")})
for r in rr["pair_rows"]: print("  ", r)
for r in rr["from_rows"]: print("  from", r)
print("HESA:", {k: v for k, v in out["hesa"].items() if k != "diff_rows"}); print(" diffs:", out["hesa"]["diff_rows"])
